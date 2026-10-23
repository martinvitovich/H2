"""Planillas y documentación del proyecto integral: tablas resumen comunes, exportación a Excel
(openpyxl) e informes PDF (ReportLab) del proyecto y de cada elemento.

Todas las salidas identifican el proyecto, la fecha, las unidades, la norma y el estado de validación,
y distinguen resultados vigentes de desactualizados o pendientes: nunca se presenta como comprobado
un elemento con cálculos pendientes.
"""

from __future__ import annotations

import io
from datetime import datetime

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, SimpleDocTemplate, Spacer

from services.report_pdf import _p, _register_fonts, _styles, _table
from version import SOFTWARE_NAME, __version__

from .analisis import CALC_LABEL
from .model import TYPE_LABEL
from .servicio import STATUS_LABEL

COMPLIANCE_LABEL = {"cumple": "Cumple", "no_cumple": "No cumple", "no_evaluable": "No evaluable", None: "—"}
LINK_LABEL = {"losa_viga": "Losa → viga", "viga_apoyo": "Viga → apoyo", "pilar_pilar": "Pilar → pilar inferior",
              "pilar_zapata": "Pilar → zapata"}
UNITS_NOTE = ("Unidades: luces y dimensiones en planta en m; secciones y cantos en mm; cargas superficiales en kN/m², "
              "lineales en kN/m, puntuales y axiles en kN; momentos en kN·m; tensiones del terreno en kPa. "
              "G: permanente característica; Q: variable característica.")


def _n(v, nd: int = 2):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return round(float(v), nd)
    return v


def _s(v, nd: int = 2) -> str:
    v = _n(v, nd)
    if v is None:
        return "—"
    if isinstance(v, float):
        s = f"{v:.{nd}f}".rstrip("0").rstrip(".")
        return s.replace(".", ",")
    return str(v)


def _layers(layers: list[dict] | None) -> str:
    if not layers:
        return ""
    by: dict[str, list[str]] = {}
    for l in layers:
        by.setdefault(l["face"], []).append(f"{l['n']}Ø{int(l['diameter'])}")
    return "; ".join(f"{k}: {' + '.join(v)}" for k, v in by.items())


def build_tables(ov: dict) -> list[dict]:
    """Tablas resumen del proyecto: [{"key", "title", "headers", "rows"}] con valores numéricos crudos."""
    els = ov["elements"]
    code = {e["id"]: e["code"] for e in els}
    lv = {l["id"]: l for l in ov["levels"]}
    by = lambda t: [e for e in els if e["type"] == t]  # noqa: E731
    st = lambda e: STATUS_LABEL.get(e["status"], e["status"])  # noqa: E731
    out_links: dict[str, list[dict]] = {}
    in_links: dict[str, list[dict]] = {}
    for l in ov["links"]:
        out_links.setdefault(l["source_id"], []).append(l)
        in_links.setdefault(l["target_id"], []).append(l)
    T = []
    T.append({"key": "niveles", "title": "Plantas y niveles", "headers": ["Código", "Nombre", "Cota (m)",
                                                                          "Altura de planta (m)", "N.º elementos"],
              "rows": [[l["code"], l["name"], _n(l["elevation_m"]), _n(l["height_m"]),
                        sum(1 for e in els if e["level_id"] == l["id"])] for l in ov["levels"]]})
    rows = []
    for e in by("losa"):
        d, L = e["data"], e["loads"]
        area = d["lx"] * d["ly"] if d.get("lx") and d.get("ly") else None
        eff = L.get("efforts") or {}
        rows.append([e["code"], e["level_code"], d.get("tipo"), _n(d.get("lx")), _n(d.get("ly")), _n(area),
                     _n(d.get("h"), 0), _n(d.get("cm")), _n(d.get("scu")), _n(d.get("muros")), _n(L.get("g")),
                     _n(L.get("q")), d.get("modelo"), _n(d.get("luz")), _n(eff.get("m_pos")), _n(eff.get("m_neg")),
                     _n(eff.get("v_ed")), ", ".join(code[x["target_id"]] for x in out_links.get(e["id"], [])), st(e)])
    T.append({"key": "losas", "title": "Losas: dimensiones y cargas",
              "headers": ["Código", "Nivel", "Tipo", "Largo (m)", "Ancho (m)", "Superficie (m²)", "h (mm)",
                          "CM (kN/m²)", "SCU (kN/m²)", "Muros (kN)", "g (kN/m²)", "q (kN/m²)", "Esfuerzos",
                          "Luz (m)", "m_Ed⁺ (kN·m/m)", "|m_Ed⁻| (kN·m/m)", "v_Ed (kN/m)", "Transmite a", "Estado"],
              "rows": rows})
    rows = []
    for e in by("viga"):
        d, L = e["data"], e["loads"]
        eff = L.get("efforts") or {}
        R = L.get("reactions") or {}
        ri, rj = R.get("inicio"), R.get("fin")
        sup = [f"{code[x['target_id']]} ({x['data'].get('extremo')})" for x in out_links.get(e["id"], [])
               if x["kind"] == "viga_apoyo"]
        rows.append([e["code"], e["level_code"], _n(d.get("b"), 0), _n(d.get("h"), 0), _n(d.get("L")),
                     d.get("sistema"), d.get("posicion"), "sí" if d.get("invertida") else "no", _n(L.get("g")),
                     _n(L.get("q")), _n(eff.get("M_pos")), _n(eff.get("M_neg")), _n(eff.get("V_ed")),
                     _n(eff.get("M_qp")), _n(ri[0]) if ri else None, _n(ri[1]) if ri else None,
                     _n(rj[0]) if rj else None, _n(rj[1]) if rj else None, ", ".join(sup), st(e)])
    T.append({"key": "vigas", "title": "Vigas: sección, cargas, esfuerzos y reacciones",
              "headers": ["Código", "Nivel", "b (mm)", "h (mm)", "L (m)", "Sistema", "Posición", "Invertida",
                          "g (kN/m)", "q (kN/m)", "M_Ed⁺ (kN·m)", "|M_Ed⁻| (kN·m)", "V_Ed (kN)", "M_qp (kN·m)",
                          "R_G inicio (kN)", "R_Q inicio (kN)", "R_G fin (kN)", "R_Q fin (kN)", "Apoya en", "Estado"],
              "rows": rows})
    rows = []
    for e in by("pilar"):
        d, L = e["data"], e["loads"]
        eff = L.get("efforts") or {}
        rec = ", ".join(code[x["source_id"]] for x in in_links.get(e["id"], []))
        sup = ", ".join(code[x["target_id"]] for x in out_links.get(e["id"], []))
        rows.append([e["code"], e["level_code"], _n(d.get("b"), 0), _n(d.get("h"), 0),
                     _n(eff.get("altura") or d.get("altura")), _n(d.get("beta")), _n(L.get("N_G")), _n(L.get("N_Q")),
                     _n(eff.get("N_Ed")), _n(d.get("M01")), _n(d.get("M02")), rec, sup, st(e)])
    T.append({"key": "pilares", "title": "Pilares: dimensiones y bajada de cargas",
              "headers": ["Código", "Nivel", "b (mm)", "h (mm)", "Altura (m)", "β", "N_G base (kN)", "N_Q base (kN)",
                          "N_Ed (kN)", "M₀₁ (kN·m)", "M₀₂ (kN·m)", "Recibe de", "Apoya en", "Estado"], "rows": rows})
    rows = []
    sadm_p = ov["project"]["settings"].get("sigma_adm")
    for e in by("zapata"):
        d, L = e["data"], e["loads"]
        pil = ", ".join(code[x["source_id"]] for x in in_links.get(e["id"], []))
        nk = (L["N_G"] or 0) + (L["N_Q"] or 0) if L.get("N_G") is not None else None
        rows.append([e["code"], _n(d.get("B")), _n(d.get("L")), _n(d.get("h"), 0), pil, _n(nk),
                     _n(d.get("sigma_adm") or sadm_p), st(e)])
    T.append({"key": "zapatas", "title": "Zapatas", "headers": ["Código", "B (m)", "L (m)", "h (mm)", "Pilar",
                                                                "N_k (kN)", "σ_adm (kPa)", "Estado"], "rows": rows})
    T.append({"key": "relaciones", "title": "Relaciones entre elementos",
              "headers": ["Tipo", "Origen", "Destino", "Datos"],
              "rows": [[LINK_LABEL[l["kind"]], code[l["source_id"]], code[l["target_id"]],
                        "; ".join(f"{k} = {_s(v)}" for k, v in l["data"].items() if v not in (None, ""))]
                       for l in ov["links"]]})
    rows = []
    for e in els:
        for ct, c in e["calcs"].items():
            if c["status"] == "no_aplica":
                continue
            ok = c.get("ok") if c["status"] == "actual" else None
            rows.append([e["code"], CALC_LABEL[ct], STATUS_LABEL.get(c["status"], c["status"]),
                         ("Cumple" if ok else "No cumple" if ok is False else "No evaluable")
                         if c["status"] == "actual" else "—",
                         (c.get("texto") or "") + (" [resultado desactualizado]" if c["status"] == "desactualizado"
                                                   and c.get("texto") else "")])
    T.append({"key": "resultados", "title": "Resultados y verificaciones",
              "headers": ["Elemento", "Cálculo", "Estado", "Resultado", "Resumen"], "rows": rows})
    rows = []
    for e in els:
        d = e["data"]
        if e["type"] == "viga":
            prop = " / ".join(filter(None, [(e["calcs"].get(ct, {}).get("summary") or {}).get("armado")
                                           for ct in ("viga_flexion", "viga_flexion_neg")
                                           if e["calcs"].get(ct, {}).get("status") == "actual"]))
            cs = e["calcs"].get("viga_cortante", {})
            est = (cs.get("summary") or {}).get("estribos") if cs.get("status") == "actual" else None
            rows.append([e["code"], "Viga", f"{_s(d.get('b'), 0)}×{_s(d.get('h'), 0)}", _layers(d.get("layers")),
                         prop if not d.get("layers") else "", est or (
                             f"cØ{_s(d.get('stirrup_phi'), 0)} c/{_s(d.get('stirrup_s'), 0)}" if d.get("stirrup_s") else "")])
        elif e["type"] == "pilar":
            rows.append([e["code"], "Pilar", f"{_s(d.get('b'), 0)}×{_s(d.get('h'), 0)}", _layers(d.get("layers")), "",
                         f"cØ{_s(d.get('stirrup_phi'), 0)} c/{_s(d.get('stirrup_s'), 0)}" if d.get("stirrup_s") else ""])
        elif e["type"] == "losa":
            dis = "; ".join(filter(None, [f"inf. Ø{_s(d.get('phi'), 0)} c/{_s(d.get('s'), 0)}" if d.get("s") else "",
                                          f"sup. Ø{_s(d.get('phi_sup'), 0)} c/{_s(d.get('s_sup'), 0)}"
                                          if d.get("s_sup") else ""]))
            prop = "; ".join(filter(None, [(e["calcs"].get(ct, {}).get("summary") or {}).get("armado")
                                           for ct in ("losa_flexion", "losa_flexion_neg")
                                           if e["calcs"].get(ct, {}).get("status") == "actual"]))
            rows.append([e["code"], "Losa", f"h = {_s(d.get('h'), 0)}", dis, prop if not dis else "", ""])
    T.append({"key": "armaduras", "title": "Resumen de armaduras",
              "headers": ["Elemento", "Tipo", "Sección (mm)", "Armadura dispuesta", "Propuesta vigente (dimensionado)",
                          "Estribos / cercos"], "rows": rows})
    rows = []
    for e in els:
        for m in e["missing"]:
            rows.append([e["code"], "Faltan datos", m])
        for m in e["errors"]:
            rows.append([e["code"], "Error", m])
        for ct, c in e["calcs"].items():
            if c["status"] in ("desactualizado", "sin_calcular", "error"):
                rows.append([e["code"], STATUS_LABEL[c["status"]], CALC_LABEL[ct]])
        for m in e["warnings"]:
            rows.append([e["code"], "Advertencia", m])
    T.append({"key": "pendientes", "title": "Pendientes, advertencias y errores",
              "headers": ["Elemento", "Asunto", "Detalle"], "rows": rows})
    return T


def _header_rows(ov: dict) -> list[list]:
    p = ov["project"]
    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    v = ov["summary"]["verification"]
    return [["Proyecto", p["name"]], ["Ubicación", p["location"] or "—"], ["Cliente", p["client"] or "—"],
            ["Fecha del proyecto", p["project_date"] or "—"], ["Sistema estructural", p["structural_system"] or "—"],
            ["Norma", f"{p['norm']['name']} — {p['norm']['edition']}"],
            ["Software", f"{SOFTWARE_NAME} v{__version__}"], ["Fecha de emisión", now],
            ["Estado de comprobación", v["text"]], ["Estado de validación del software", v["disclaimer"]],
            ["Unidades", UNITS_NOTE]]


def build_excel(ov: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Proyecto"
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E8EEF5")
    ws.append(["HA-Anejo19 — Planilla del proyecto"])
    ws["A1"].font = Font(bold=True, size=14)
    for r in _header_rows(ov):
        ws.append(r)
        ws.cell(ws.max_row, 1).font = bold
    ws.append([])
    ws.append(["Configuración del proyecto", "Valor"])
    ws.cell(ws.max_row, 1).font = bold
    for k, v in ov["project"]["settings"].items():
        ws.append([k, v])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=2, max_col=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    for t in build_tables(ov):
        sh = wb.create_sheet(t["title"][:31].replace(":", "").replace("/", "-"))
        sh.append([f"{ov['project']['name']} — {t['title']}"])
        sh["A1"].font = Font(bold=True, size=12)
        sh.append([f"{SOFTWARE_NAME} v{__version__} · {ov['project']['norm']['name']} · "
                   f"{datetime.now().strftime('%d/%m/%Y %H:%M')} · Validación normativa del software pendiente"])
        sh.append([])
        sh.append(t["headers"])
        for i in range(1, len(t["headers"]) + 1):
            c = sh.cell(4, i)
            c.font = bold
            c.fill = fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
            sh.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(t["headers"][i - 1]) + 4))
        for r in t["rows"]:
            sh.append(["" if v is None else v for v in r])
        sh.freeze_panes = "B5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _doc(buf, title: str, project: str, landscape_mode: bool = False):
    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    size = landscape(A4) if landscape_mode else A4

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("DejaVu", 7)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(12 * mm, 8 * mm, f"{SOFTWARE_NAME} v{__version__} · Código Estructural, Anejo 19 · {now} · "
                                           "Validación normativa del software pendiente")
        canvas.drawRightString(size[0] - 12 * mm, 8 * mm, f"Página {doc.page}")
        canvas.drawString(12 * mm, size[1] - 9 * mm, f"{project} — {title}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=size, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=14 * mm,
                            bottomMargin=14 * mm, title=f"{title} — {project}", author=SOFTWARE_NAME)
    return doc, on_page, size[0] - 24 * mm


def _fmt_row(r: list) -> list[str]:
    return [_s(v) if not isinstance(v, str) else v for v in r]


def build_project_pdf(ov: dict, history: list[dict]) -> bytes:
    _register_fonts()
    S = _styles()
    buf = io.BytesIO()
    doc, on_page, W = _doc(buf, "Memoria resumen del proyecto", ov["project"]["name"], landscape_mode=True)
    E: list = [_p("Memoria resumen del proyecto estructural", S["h1"])]
    E.append(_table([[a, b] for a, b in _header_rows(ov)], [55 * mm, W - 55 * mm], S["base"], header=False))
    E.append(Spacer(1, 4))
    E.append(_p("AVISO: herramienta de apoyo académico y profesional. Los resultados deben ser revisados "
                "independientemente por un ingeniero estructural responsable. Las transmisiones de cargas usan "
                "equilibrio de vigas isostáticas y áreas tributarias definidas por el usuario (sin continuidad ni "
                "alternancia de sobrecargas); los valores por defecto de coeficientes de acciones deben confirmarse.",
                S["warn"]))
    sm = ov["summary"]
    E.append(_p("Estado del proyecto", S["h2"]))
    E.append(_table([["Estado", "N.º de elementos"]] +
                    [[STATUS_LABEL.get(k, k), str(v)] for k, v in sm["by_status"].items()] +
                    [[f"Resultado: {COMPLIANCE_LABEL.get(k, k)}", str(v)] for k, v in sm["by_compliance"].items()],
                    [80 * mm, 40 * mm], S["small"]))
    E.append(_p("Configuración y materiales", S["h2"]))
    st = ov["project"]["settings"]
    E.append(_table([["Dato", "Valor"]] + [[k, _s(v) if not isinstance(v, str) else v] for k, v in st.items()],
                    [80 * mm, 60 * mm], S["small"]))
    if ov["project"]["parameter_overrides"]:
        E.append(_p("Parámetros normativos modificados por el usuario: " +
                    "; ".join(f"{k} = {v}" for k, v in ov["project"]["parameter_overrides"].items()), S["warn"]))
    for t in build_tables(ov):
        E.append(_p(t["title"], S["h2"]))
        if not t["rows"]:
            E.append(_p("Sin datos.", S["small"]))
            continue
        n = len(t["headers"])
        widths = [W / n] * n
        if t["key"] in ("resultados", "pendientes", "relaciones", "armaduras"):
            widths = {"resultados": [22, 60, 30, 25], "pendientes": [25, 35], "relaciones": [40, 30, 30],
                      "armaduras": [22, 14, 28, 70, 70]}[t["key"]]
            widths = [w * mm for w in widths]
            widths.append(W - sum(widths))
        E.append(_table([t["headers"]] + [_fmt_row(r) for r in t["rows"]], widths, S["small"]))
    E.append(PageBreak())
    E.append(_p("Registro de cambios (últimas 60 entradas)", S["h2"]))
    E.append(_table([["Fecha (UTC)", "Elemento", "Acción", "Detalle"]] +
                    [[h["ts"].replace("T", " ")[:19], h["element_code"] or "—", h["action"].replace("_", " "),
                      _hist_text(h["detail"])] for h in history[:60]],
                    [35 * mm, 25 * mm, 40 * mm, W - 100 * mm], S["small"]))
    doc.build(E, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


def _hist_text(d: dict) -> str:
    if "cambios" in d:
        return "; ".join(f"{k}: {_s(v.get('antes')) if not isinstance(v.get('antes'), (list, dict)) else '…'} → "
                         f"{_s(v.get('despues')) if not isinstance(v.get('despues'), (list, dict)) else '…'}"
                         for k, v in d["cambios"].items())[:400]
    return "; ".join(f"{k}: {v}" for k, v in d.items() if not isinstance(v, (dict, list)))[:400]


def build_element_pdf(ov: dict, det: dict, history: list[dict]) -> bytes:
    """Ficha del elemento: identificación, datos, relaciones, procedencia de cargas, estado y verificaciones."""
    _register_fonts()
    S = _styles()
    buf = io.BytesIO()
    title = f"Ficha de {TYPE_LABEL[det['type']].lower()} {det['code']}"
    doc, on_page, W = _doc(buf, title, ov["project"]["name"])
    E: list = [_p(title, S["h1"])]
    E.append(_table([[a, b] for a, b in _header_rows(ov)[:8]] +
                    [["Nivel", det["level_code"] or "—"], ["Descripción", det["description"] or "—"],
                     ["Estado del elemento", STATUS_LABEL.get(det["status"], det["status"]) +
                      (f" · {COMPLIANCE_LABEL.get(det['compliance'])}" if det["compliance"] else "")]],
                    [45 * mm, W - 45 * mm], S["base"], header=False))
    E.append(Spacer(1, 3))
    E.append(_p("AVISO: los resultados deben ser revisados independientemente por un ingeniero estructural "
                "responsable.", S["warn"]))
    E.append(_p("1. Datos de entrada", S["h2"]))
    rows = [["Dato", "Valor"]]
    for k, v in det["data"].items():
        if isinstance(v, list):
            v = _layers(v) if k == "layers" else "; ".join(
                f"x = {_s(p['pos'])} m: G = {_s(p['G'])}, Q = {_s(p['Q'])} kN" for p in v)
        rows.append([k, _s(v) if not isinstance(v, str) else v])
    E.append(_table(rows, [60 * mm, W - 60 * mm], S["small"]))
    E.append(_p("2. Relaciones", S["h2"]))
    E.append(_table([["Sentido", "Tipo", "Elemento", "Datos"]] +
                    [[r["direction"], LINK_LABEL[r["kind"]], r["other_code"],
                      "; ".join(f"{k} = {_s(v)}" for k, v in r["data"].items() if v not in (None, ""))]
                     for r in det["relations"]], [22 * mm, 40 * mm, 25 * mm, W - 87 * mm], S["small"]))
    if det["affects"]:
        E.append(_p("Elementos que dependen de este: " + ", ".join(det["affects"]), S["small"]))
    E.append(_p("3. Cargas y procedencia", S["h2"]))
    ld = det["loads_detail"]
    E.append(_table([["Concepto", "Expresión", "G", "Q", "Unidad", "Origen"]] +
                    [[r["concepto"], r["expresion"], _s(r["G"]), _s(r["Q"]), r["unidad"], r["origen"]]
                     for r in ld["rows"]], [35 * mm, 62 * mm, 16 * mm, 16 * mm, 22 * mm, W - 151 * mm], S["small"]))
    for m in ld["missing"]:
        E.append(_p("Faltan datos: " + m, S["warn"]))
    for m in ld["errors"] + ld["warnings"]:
        E.append(_p(m, S["warn"]))
    for m in ld["notes"]:
        E.append(_p(m, S["small"]))
    E.append(_p("4. Cálculos y verificaciones", S["h2"]))
    for c in det["calc_list"]:
        if c["status"] == "no_aplica":
            continue
        stored = c.get("stored")
        head = f"{c['label']} — {STATUS_LABEL.get(c['status'], c['status'])}"
        if stored:
            head += f" (calculado {stored['computed_at'].replace('T', ' ')[:16]} UTC, motor v{stored['engine_version']})"
        E.append(_p(head, S["base"]))
        if c["status"] == "faltan_datos":
            E.append(_p("Faltan datos: " + "; ".join(c.get("missing", [])), S["warn"]))
            continue
        if c["status"] != "actual":
            E.append(_p(c.get("reason") or "Resultado no vigente: recalcule el elemento.", S["warn"]))
            continue
        out = stored["output"]
        E.append(_p(out["summary"].get("texto") or "", S["small"]))
        tr = out["result"]["trace"]
        E.append(_table([["Verificación", "Valor", "Límite / capacidad", "Resultado", "Referencia"]] +
                        [[k["name"], k["demand"], k["capacity"],
                          "CUMPLE" if k["ok"] else "NO CUMPLE" if k["ok"] is False else "NO EVALUABLE", k["reference"]]
                         for k in tr["checks"]], [50 * mm, 35 * mm, 35 * mm, 22 * mm, W - 142 * mm], S["small"]))
        for w in tr["warnings"]:
            E.append(_p("Advertencia: " + w, S["warn"]))
        E.append(Spacer(1, 4))
    E.append(_p("5. Historial del elemento", S["h2"]))
    E.append(_table([["Fecha (UTC)", "Acción", "Detalle"]] +
                    [[h["ts"].replace("T", " ")[:19], h["action"].replace("_", " "), _hist_text(h["detail"])]
                     for h in history[:40]], [35 * mm, 35 * mm, W - 70 * mm], S["small"]))
    doc.build(E, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()
